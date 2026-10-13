from io import BytesIO
from pathlib import Path

def extract_text(file_storage, filename):
    ext = Path(filename).suffix.lower()
    data = file_storage.read()
    if ext == ".txt":
        return data.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        import fitz  # PyMuPDF
        doc = fitz.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)
    if ext == ".docx":
        from docx import Document
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Unsupported file type")
